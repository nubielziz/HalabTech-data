"""Build the Round-7 workbook: all Rounds 1-6 sheets preserved untouched + new R7 sheets placed first.

Usage: python3 build_r7.py <r6_workbook.xlsx> <r7_sheets_spec.json> <out.xlsx>
spec = {"sheets": [{"title", "header", "rows", "widths"?, "note"?, "conf_cols"?, "url_cols"?, "freeze"?}], "meta": {...}}
"""
import json, re, sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

R6, SPEC, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
spec = json.load(open(SPEC))
wb = load_workbook(R6)
r6_titles = list(wb.sheetnames)

HDR_FILL = PatternFill('solid', fgColor='0B3C5D')
HDR_FONT = Font(bold=True, color='FFFFFF')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CONF = {'CONFIRMED': 'C6EFCE', 'STRONG': 'DDEBF7', 'LIKELY': 'FFF2CC', 'UNKNOWN': 'FCE4D6', 'DISPUTED': 'F4B6C2', 'REJECTED': 'D9D9D9'}
existing_tables = {t for ws in wb.worksheets for t in ws.tables.keys()}


def fill_for(v):
    if not isinstance(v, str):
        return None
    u = v.upper()
    for k, c in CONF.items():
        if u.startswith(k):
            return PatternFill('solid', fgColor=c)
    return None


def add_sheet(s, index):
    title = s['title'][:31]
    ws = wb.create_sheet(title, index)
    start = 1
    if s.get('note'):
        ws.cell(row=1, column=1, value=s['note']).font = Font(italic=True, color='404040')
        start = 2
    hdr = s['header']
    for ci, h in enumerate(hdr, 1):
        c = ws.cell(row=start, column=ci, value=h)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center')
    conf_cols = set(s.get('conf_cols', []))
    url_cols = set(s.get('url_cols', []))
    for ri, r in enumerate(s['rows'], start + 1):
        for ci, h in enumerate(hdr, 1):
            v = r.get(h, '')
            if isinstance(v, (list, tuple)):
                v = '; '.join(str(x) for x in v)
            if isinstance(v, dict):
                v = json.dumps(v, ensure_ascii=False)
            c = ws.cell(row=ri, column=ci, value=v)
            c.alignment, c.border = WRAP, BORDER
            if h in conf_cols:
                f = fill_for(v)
                if f:
                    c.fill = f
            if h in url_cols and isinstance(v, str) and v.startswith('http') and ' ' not in v:
                c.hyperlink = v
                c.font = Font(color='0563C1', underline='single')
    widths = s.get('widths', {})
    for ci, h in enumerate(hdr, 1):
        w = widths.get(h)
        if w is None:
            sample = [len(str(h))] + [len(str(r.get(h, ''))) for r in s['rows'][:300]]
            sample.sort()
            w = min(max(10, int(sample[int(len(sample) * 0.9)] * 0.9) + 2), 60)
        ws.column_dimensions[get_column_letter(ci)].width = w
    if s.get('freeze', True):
        ws.freeze_panes = ws.cell(row=start + 1, column=2)
    if s['rows']:
        name = 'R7' + re.sub(r'[^A-Za-z0-9]', '', title)[:26]
        while name in existing_tables:
            name += 'x'
        existing_tables.add(name)
        t = Table(displayName=name, ref=f"A{start}:{get_column_letter(len(hdr))}{start + len(s['rows'])}")
        t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
        ws.add_table(t)
    return ws


for i, s in enumerate(spec['sheets']):
    add_sheet(s, i)
wb.active = 0
wb.save(OUT)
print('saved', OUT)
print('new sheets:', [s['title'] for s in spec['sheets']])
print('preserved R1-R6 sheets:', r6_titles)
